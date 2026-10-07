import io

import httpx
from openpyxl import load_workbook

from arte.crawler import run_job
from arte.fetch.http import Fetcher
from arte.jobs import Job, NotifyFilter
from tests.conftest import fixture

SEARCH = """<html><body>
<a href="/sale/flat/312345678/">1</a><a href="/sale/flat/398765432/">2</a></body></html>"""


def make_transport(state):
    def handler(request: httpx.Request):
        path = request.url.path
        if path == "/cat.php":
            page = request.url.params.get("p", "1")
            return httpx.Response(200, text=SEARCH if page == "1" else "<html></html>")
        if path == "/sale/flat/312345678/":
            html = fixture("cian_like.html").replace("21900000", str(state["price"]))
            return httpx.Response(200, text=html)
        if path == "/sale/flat/398765432/":
            return httpx.Response(404 if state.get("gone") else 200, text=fixture("cian_like.html")
                                  .replace("312345678", "398765432").replace("78.4", "61"))
        return httpx.Response(404)
    return httpx.MockTransport(handler)


async def test_full_crawl_cycle_with_api(database, settings, monkeypatch):
    state = {"price": 21_900_000}
    job = Job(name="msk", source="cian", search_urls=["https://www.cian.ru/cat.php?deal_type=sale"],
              max_pages=3, notify=NotifyFilter(kinds=["new", "price_down"], price_max=25_000_000))
    sent = []

    async def fake_notify(events, job_, settings_):
        sent.extend(e.kind for e in events)

    monkeypatch.setattr("arte.notify.telegram.notify_events", fake_notify)

    async with Fetcher(settings, transport=make_transport(state)) as f:
        stats = await run_job(job, settings, fetcher=f)
    assert (stats.search_pages, stats.listing_urls, stats.saved, stats.errors) == (2, 2, 2, 0)
    assert sent == ["new", "new"]

    state.update(price=20_500_000, gone=True)
    async with Fetcher(settings, transport=make_transport(state)) as f:
        stats = await run_job(job, settings, fetcher=f)
    assert stats.removed == 1
    assert sorted(e.kind for e in stats.events) == ["price_down", "removed"]

    from arte.api.app import app
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t") as client:
        page = (await client.get("/api/listings")).json()
        assert page["total"] == 1
        item = page["items"][0]
        assert item["price"] == 20_500_000 and item["source"] == "cian"
        detail = (await client.get(f"/api/listings/{item['id']}")).json()
        assert [p["price"] for p in detail["price_history"]] == [21_900_000, 20_500_000]
        assert (await client.get("/api/listings", params={"active": "false"})).json()["total"] == 1
        assert (await client.get("/api/listings", params={"price_max": 1_000_000})).json()["total"] == 0
        stats = (await client.get("/api/stats")).json()
        assert stats["active_by_source"] == {"cian": 1}
        assert stats["price_changes"] == 1
        assert {e["kind"] for e in (await client.get("/api/events")).json()} == {"new", "price_down", "removed"}

        xlsx = await client.get("/api/export.xlsx")
        ws = load_workbook(io.BytesIO(xlsx.content)).active
        assert ws["A1"].value == "ID" and ws.max_row == 2
        csv_text = (await client.get("/api/export.csv")).content.decode("utf-8-sig")
        assert csv_text.splitlines()[0].startswith("ID;Площадка")
        assert (await client.get("/")).status_code == 200


def test_notify_filter():
    f = NotifyFilter(kinds=["new"], price_max=10, rooms=[1, 2])
    assert f.accepts("new", 9, 1, None)
    assert not f.accepts("new", 11, 1, None)
    assert not f.accepts("new", 9, 3, None)
    assert not f.accepts("price_up", 9, 1, None)
