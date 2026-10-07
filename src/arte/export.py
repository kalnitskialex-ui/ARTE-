"""Выгрузка в CSV / Excel / JSON."""

import csv
import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from arte.storage.queries import EXPORT_COLUMNS

HEADERS_RU = {
    "id": "ID", "source": "Площадка", "external_id": "ID на площадке", "url": "Ссылка",
    "title": "Заголовок", "deal_type": "Сделка", "property_type": "Тип", "price": "Цена",
    "currency": "Валюта", "price_per_m2": "Цена за м²", "area_total": "Площадь",
    "area_living": "Жилая", "area_kitchen": "Кухня", "rooms": "Комнат", "floor": "Этаж",
    "floors_total": "Этажей", "year_built": "Год постройки", "address": "Адрес", "metro": "Метро",
    "lat": "Широта", "lon": "Долгота", "seller_type": "Продавец", "object_id": "ID объекта",
    "is_active": "Активно", "first_seen": "Впервые", "last_seen": "Последний раз",
}


def _plain(v):
    return v.isoformat(sep=" ", timespec="seconds") if isinstance(v, datetime) else v


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS, delimiter=";")
    writer.writerow(HEADERS_RU)
    for r in rows:
        writer.writerow({k: _plain(v) for k, v in r.items()})
    # BOM — чтобы Excel сразу открыл кириллицу правильно.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_json(rows: list[dict]) -> bytes:
    return json.dumps([{k: _plain(v) for k, v in r.items()} for r in rows],
                      ensure_ascii=False, indent=2).encode("utf-8")


def to_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Объявления"
    ws.append([HEADERS_RU[c] for c in EXPORT_COLUMNS])
    for r in rows:
        ws.append([r.get(c) for c in EXPORT_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"url": 40, "title": 40, "address": 45}
    for i, col in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 13)
    for row in ws.iter_rows(min_row=2):
        for cell, col in zip(row, EXPORT_COLUMNS, strict=True):
            if col in ("price", "price_per_m2"):
                cell.number_format = "# ##0"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FORMATS = {
    "csv": (to_csv, "text/csv; charset=utf-8"),
    "json": (to_json, "application/json"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}
